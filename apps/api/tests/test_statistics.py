"""FR-24~FR-26 / AC-19·AC-20 — 학생·반 통계와 엑셀 내보내기."""
import asyncio

from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

WEEKS = ["2026-09-04", "2026-09-11", "2026-09-18"]  # 3주 연속 금요일


def _write(api, class_id, on, records, notes=None):
    session_id = api.get(
        "/api/daily", params={"class_id": class_id, "date": on}
    ).json()["session"]["id"]
    if notes:
        api.put(f"/api/daily/{session_id}/notes", json=notes)
    api.put(f"/api/daily/{session_id}/records", json={"records": records})


def _history(api, klass):
    """3주간 학생1은 85·90·95점 A/A/B, 학생2는 65·70·75점 B/C/C. 학생3은 마지막 주 결석."""
    first, second, third = klass["student_ids"]
    scores = {first: [85, 90, 95], second: [65, 70, 75], third: [50, 55, None]}
    grades = {first: ["A", "A", "B"], second: ["B", "C", "C"], third: ["C", "D", None]}
    for index, on in enumerate(WEEKS):
        records = []
        for student_id in (first, second, third):
            score = scores[student_id][index]
            records.append(
                {
                    "student_id": student_id,
                    "attendance_status": "present" if score is not None else "absent",
                    "homework_grade": grades[student_id][index],
                    "test_score_num": score,
                }
            )
        _write(api, klass["class_id"], on, records, {"test_name": "주간", "test_max_score": 100})


def _exam(api, migrated_database, class_id, on, scores, name="9월 모의"):
    """OMR 채점 시험. 판독·채점은 TASK-34~36에서 검사했으므로 응시 점수를 직접 넣는다."""
    from mathdesk.models import ExamAttempt

    exam_id = api.post(
        "/api/exams", json={"name": name, "class_id": class_id, "exam_date": on}
    ).json()["id"]

    async def run():
        engine = create_async_engine(migrated_database)
        async with async_sessionmaker(engine)() as session:
            for student_id, score in scores.items():
                session.add(
                    ExamAttempt(exam_id=exam_id, student_id=student_id, score=score, source="omr")
                )
            await session.commit()
        await engine.dispose()

    asyncio.run(run())
    return exam_id


def _rounds(api, student_id, **params):
    return api.get(f"/api/stats/students/{student_id}", params={"to": WEEKS[-1], **params})


def test_student_rounds_count_each_daily_test_and_omr_exam_with_the_class_average(
    api, klass, migrated_database
):
    """DCR-018 / AC-19: 테스트 점수 1개가 1회차 — 같은 날 일일 테스트와 OMR 시험은 2회차다."""
    _history(api, klass)
    first, second, third = klass["student_ids"]
    _exam(api, migrated_database, klass["class_id"], WEEKS[1], {first: 80, second: 60})
    # 회차가 아닌 것: 테스트 없는 수업, 종료일 뒤 시험, 시험일 없는 시험, 채점 안 된 시험
    _write(api, klass["class_id"], "2026-09-16", [], {"homework": "교재 10p"})
    _exam(api, migrated_database, klass["class_id"], "2026-09-25", {first: 100}, "10월 모의")
    _exam(api, migrated_database, klass["class_id"], None, {first: 100}, "날짜 없음")
    _exam(api, migrated_database, klass["class_id"], WEEKS[0], {}, "미채점")

    body = _rounds(api, first).json()

    assert body["student"]["name"] == "김나윤"
    assert body["total"] == 4
    assert [
        (r["date"], r["source"], r["name"], r["score"], r["class_average"]) for r in body["rounds"]
    ] == [
        ("2026-09-04", "daily", "주간", 85.0, 66.7),
        ("2026-09-11", "daily", "주간", 90.0, 71.7),
        ("2026-09-11", "exam", "9월 모의", 80.0, 70.0),
        ("2026-09-18", "daily", "주간", 95.0, 85.0),
    ]
    # 회차 날짜의 그 반 수업 기록
    assert [(r["homework_grade"], r["attendance_status"]) for r in body["rounds"]] == [
        ("A", "present"), ("A", "present"), ("A", "present"), ("B", "present")
    ]
    # 응시하지 않은 회차는 학생 점수만 빈다
    absent = _rounds(api, third).json()["rounds"]
    assert [(r["score"], r["class_average"]) for r in absent[2:]] == [(None, 70.0), (None, 85.0)]
    assert absent[3]["attendance_status"] == "absent"


def test_student_rounds_show_the_latest_ten_or_twenty(api, klass):
    student_id = klass["student_ids"][0]
    days = [f"2026-08-{day:02d}" for day in range(1, 13)]
    for index, on in enumerate(days):
        _write(api, klass["class_id"], on, [{"student_id": student_id, "test_score_num": index}])

    ten = _rounds(api, student_id).json()
    twenty = _rounds(api, student_id, rounds=20).json()

    assert ten["total"] == twenty["total"] == 12
    assert [r["date"] for r in ten["rounds"]] == days[2:]
    assert [r["date"] for r in twenty["rounds"]] == days
    assert _rounds(api, student_id, rounds=15).status_code == 422


def test_class_stats_gives_rows_and_a_period_comparison(api, klass):
    _history(api, klass)

    body = api.get(
        f"/api/stats/classes/{klass['class_id']}",
        params={
            "start": WEEKS[1],
            "end": WEEKS[2],
            "compare_start": WEEKS[0],
            "compare_end": WEEKS[0],
        },
    ).json()

    rows = {row["name"]: row for row in body["period"]["students"]}
    assert rows["김나윤"]["test_average"] == 92.5
    assert rows["김나윤"]["attendance_rate"] == 100.0
    assert rows["김나윤"]["homework_completion"] == 100.0
    assert rows["김태호"]["attendance_rate"] == 50.0
    assert rows["김정현"]["homework_completion"] == 0.0  # C·C는 기준(B) 미만
    assert body["period"]["test"]["average"] == 77.0  # 90·70·55·95·75
    assert "distribution" not in body["period"]["test"]  # DCR-018: 점수 분포 없음
    assert body["period"]["homework"]["distribution"] == {"A": 1, "B": 1, "C": 2, "D": 1}
    assert body["compare"]["test"]["average"] == 66.7  # 85·65·50
    assert body["compare"]["start"] == WEEKS[0]


def test_a_day_without_attendance_input_is_not_counted_as_a_class(api, klass):
    """TASK-86: 일일 입력에서 날짜를 열기만 해도 수업 행이 생긴다. 출결을 입력하지 않은 수업은
    출결률 분모에 넣지 않는다(열어 본 날·진도만 적은 날이 결석으로 잡히던 버그)."""
    _history(api, klass)
    api.get("/api/daily", params={"class_id": klass["class_id"], "date": "2026-09-15"})
    _write(api, klass["class_id"], "2026-09-16", [], {"homework": "교재 10p"})

    rows = {
        row["name"]: row
        for row in api.get(
            f"/api/stats/classes/{klass['class_id']}",
            params={"start": WEEKS[1], "end": WEEKS[2]},
        ).json()["period"]["students"]
    }

    assert rows["김나윤"]["attendance_rate"] == 100.0
    assert rows["김태호"]["attendance_rate"] == 50.0


def test_export_contains_the_same_rows_as_the_class_stats_screen(api, klass):
    from io import BytesIO
    from openpyxl import load_workbook

    _history(api, klass)
    params = {"class_id": klass["class_id"], "start": WEEKS[0], "end": WEEKS[2]}
    shown = api.get(f"/api/stats/classes/{klass['class_id']}", params={
        "start": WEEKS[0], "end": WEEKS[2]
    }).json()["period"]["students"]

    exported = api.get("/api/stats/export", params=params)

    assert exported.status_code == 200
    assert "spreadsheetml" in exported.headers["content-type"]
    sheet = load_workbook(BytesIO(exported.content)).active
    header, *rows = [[cell.value for cell in row] for row in sheet.iter_rows()]
    assert header == ["학생", "테스트 평균", "과제 완수율", "출결률"]
    assert rows == [
        [row["name"], row["test_average"], row["homework_completion"], row["attendance_rate"]]
        for row in shown
    ]


def test_teacher_cannot_read_history_of_a_student_outside_their_classes(api, klass):
    """설계의 권한 표: 통계는 원장이 캠퍼스 전체, 강사는 담당 반 한정이다."""
    outsider = klass["student_ids"][0]
    api.post("/api/auth/logout")
    api.sign_in("teacher_a")  # 담당 반이 없는 강사

    assert api.get(f"/api/stats/students/{outsider}").status_code == 403
